import os
import json
from pathlib import Path
from typing import Tuple, List
from .pii_detector import detect_pii, redact_text, PIIMatch


def extract_text_from_file(file_path: str) -> Tuple[str, str]:
    """Returns (text, file_type)"""
    ext = Path(file_path).suffix.lower()

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read(), "txt"

    if ext == ".pdf":
        return _extract_pdf(file_path), "pdf"

    if ext in (".docx", ".doc"):
        return _extract_docx(file_path), "docx"

    if ext in (".xlsx", ".xls"):
        return _extract_xlsx(file_path), "xlsx"

    # Fallback: try reading as text
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read(), "txt"
    except Exception:
        return "", "unknown"


def _extract_pdf(path: str) -> str:
    try:
        from pdfminer.high_level import extract_text
        return extract_text(path) or ""
    except Exception as e:
        return f"[PDF extraction failed: {e}]"


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        return f"[DOCX extraction failed: {e}]"


def _extract_xlsx(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join(str(c) if c is not None else "" for c in row)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    except Exception as e:
        return f"[XLSX extraction failed: {e}]"


def process_document(
    file_path: str,
    output_dir: str,
    custom_terms: List[str] = None,
) -> dict:
    text, file_type = extract_text_from_file(file_path)
    if not text:
        return {"error": "Could not extract text from document", "pii_count": 0, "summary": {}}

    matches, summary = detect_pii(text, custom_terms)
    redacted = redact_text(text, matches)

    base = Path(file_path).stem
    out_path = os.path.join(output_dir, f"{base}_redacted.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(redacted)

    # Also write highlighted version (wraps PII in markers for display)
    highlighted = _build_highlighted(text, matches)
    highlight_path = os.path.join(output_dir, f"{base}_highlighted.json")
    with open(highlight_path, "w") as f:
        json.dump({
            "segments": highlighted,
            "summary": summary,
            "total": len(matches),
        }, f)

    return {
        "redacted_path": out_path,
        "highlight_path": highlight_path,
        "pii_count": len(matches),
        "summary": summary,
        "file_type": file_type,
        "original_chars": len(text),
    }


def _build_highlighted(text: str, matches: List[PIIMatch]) -> list:
    if not matches:
        return [{"text": text, "pii": False}]

    sorted_matches = sorted(matches, key=lambda m: m.start)
    segments = []
    cursor = 0
    for m in sorted_matches:
        if m.start > cursor:
            segments.append({"text": text[cursor:m.start], "pii": False})
        segments.append({"text": m.text, "pii": True, "label": m.label, "replacement": m.replacement})
        cursor = m.end
    if cursor < len(text):
        segments.append({"text": text[cursor:], "pii": False})
    return segments
